"""DOCX parsing for the experimental hybrid target; independent of DocumentContent."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Length
from docx.table import Table
from docx.text.paragraph import Paragraph

from .hybrid_model import (
    HybridCell, HybridDocumentPlan, HybridListSpec, HybridParagraph, HybridRun, HybridTable, HybridUnsupported,
)


class HybridInputError(ValueError):
    pass


_BUILTIN_STYLES = {"Normal", *(f"Heading {number}" for number in range(1, 10))}


@dataclass(slots=True)
class _ListGroupTracker:
    next_group_id: int = 0
    active_kind: str | None = None

    def spec_for(self, paragraph: Paragraph) -> HybridListSpec | None:
        kind = _list_kind(paragraph)
        if kind is None:
            self.break_group()
            return None
        if kind != self.active_kind:
            self.next_group_id += 1
            self.active_kind = kind
        return HybridListSpec(kind, self.next_group_id)

    def break_group(self) -> None:
        self.active_kind = None


def _measurement_to_points(value) -> float | None:
    """Convert python-docx Length values to Word's point units."""
    if value is None:
        return None
    if isinstance(value, Length):
        return float(value.pt)
    if isinstance(value, (int, float)):
        return float(value)
    return None


def load_hybrid_plan(path: Path) -> HybridDocumentPlan:
    source = path.expanduser()
    if not source.is_file():
        raise HybridInputError(f"Hybrid mode requires an existing .docx file: {source}")
    if source.suffix.lower() != ".docx":
        raise HybridInputError("Hybrid mode accepts an existing .docx file only; TXT and plain text are unsupported.")
    try:
        document = Document(BytesIO(source.read_bytes()))
    except Exception as exc:
        raise HybridInputError(f"Could not read hybrid DOCX input: {source}") from exc

    unsupported: list[HybridUnsupported] = []
    if len(document.inline_shapes):
        unsupported.append(HybridUnsupported("images", "inline images are not supported"))
    if any(paragraph.text for section in document.sections for paragraph in section.header.paragraphs + section.footer.paragraphs):
        unsupported.append(HybridUnsupported("headers/footers", "non-empty headers or footers are not supported"))

    blocks = []
    list_groups = _ListGroupTracker()
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            blocks.append(_paragraph(paragraph, list_groups.spec_for(paragraph)))
        elif child.tag == qn("w:tbl"):
            list_groups.break_group()
            table, findings = _table(Table(child, document), list_groups)
            unsupported.extend(findings)
            if table is not None:
                blocks.append(table)
            list_groups.break_group()
    return HybridDocumentPlan(source.resolve(), tuple(blocks), tuple(unsupported))


def _paragraph(paragraph: Paragraph, list_spec: HybridListSpec | None = None) -> HybridParagraph:
    style_name = getattr(paragraph.style, "name", None)
    if style_name not in _BUILTIN_STYLES:
        style_name = None
    fmt = paragraph.paragraph_format
    runs = tuple(
        HybridRun(run.text or "", bool(run.bold), bool(run.italic), bool(run.underline))
        for run in paragraph.runs
        if run.text
    )
    return HybridParagraph(
        runs=runs,
        style_name=style_name,
        alignment=int(paragraph.alignment) if paragraph.alignment is not None else None,
        left_indent=_measurement_to_points(fmt.left_indent),
        first_line_indent=_measurement_to_points(fmt.first_line_indent),
        line_spacing=_measurement_to_points(fmt.line_spacing) if isinstance(fmt.line_spacing, Length) else None,
        list_spec=list_spec,
    )


def _list_kind(paragraph: Paragraph) -> str | None:
    name = getattr(paragraph.style, "name", "").lower()
    if name.startswith("list bullet"):
        return "bullet"
    if name.startswith("list number"):
        return "number"
    return None


def _table(table: Table, list_groups: _ListGroupTracker) -> tuple[HybridTable | None, list[HybridUnsupported]]:
    findings: list[HybridUnsupported] = []
    rows: list[tuple[HybridCell, ...]] = []
    width = None
    for row in table.rows:
        current: list[HybridCell] = []
        if width is None:
            width = len(row.cells)
        elif width != len(row.cells):
            findings.append(HybridUnsupported("complex table", "non-rectangular table grid"))
        for cell in row.cells:
            properties = cell._tc.tcPr
            if properties is not None and (
                properties.find(qn("w:gridSpan")) is not None or properties.find(qn("w:vMerge")) is not None
            ):
                findings.append(HybridUnsupported("merged table", "merged cells are not supported"))
                continue
            if cell.tables:
                findings.append(HybridUnsupported("nested table", "nested tables are not supported"))
            list_groups.break_group()
            paragraphs = tuple(_paragraph(item, list_groups.spec_for(item)) for item in cell.paragraphs)
            list_groups.break_group()
            if len(paragraphs) > 1:
                findings.append(HybridUnsupported("multi-paragraph cell", "only one paragraph per table cell is supported"))
            current.append(HybridCell(paragraphs))
        rows.append(tuple(current))
    if findings:
        return None, findings
    widths = tuple(_measurement_to_points(cell.width) for cell in table.rows[0].cells) if table.rows else ()
    return HybridTable(tuple(rows), widths), findings
