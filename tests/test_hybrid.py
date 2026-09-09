from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.dml import MSO_THEME_COLOR_INDEX
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor

import autotype.cli as cli_module
from autotype.actions import TypeText
from autotype.config import TypingConfig
from autotype.controller import RunResult, RunState
from autotype.focus import FocusError, WordFocusGuard
from autotype.hybrid_input import load_hybrid_plan
from autotype.hybrid_model import HybridCell, HybridDocumentPlan, HybridListSpec, HybridParagraph, HybridRun, HybridTable, HybridUnsupported
from autotype.hybrid_runner import HybridRunError, HybridRunner
from autotype.hybrid_word import HybridWordAdapter, HybridWordError


def _docx(tmp_path: Path) -> Path:
    path = tmp_path / "hybrid.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.alignment = 1
    paragraph.paragraph_format.left_indent = 1440
    paragraph.add_run("Bold").bold = True
    paragraph.add_run(" italic").italic = True
    doc.add_paragraph("Bullet", style="List Bullet")
    doc.add_paragraph("Number", style="List Number")
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("A1", "B1"), ("A2", "B2"))):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.save(path)
    return path


def test_hybrid_parser_extracts_supported_document_features(tmp_path: Path) -> None:
    plan = load_hybrid_plan(_docx(tmp_path))
    assert plan.paragraph_count == 3
    assert plan.table_count == 1
    assert plan.cell_count == 4
    assert plan.run_count >= 4
    assert plan.list_count == 2
    assert not plan.unsupported


def test_hybrid_parser_converts_docx_measurements_to_points(tmp_path: Path) -> None:
    path = tmp_path / "measurements.docx"
    doc = Document()
    exact = doc.add_paragraph()
    exact.paragraph_format.left_indent = Inches(1)
    exact.paragraph_format.first_line_indent = Pt(18)
    exact.paragraph_format.line_spacing = Pt(12)
    relative = doc.add_paragraph()
    relative.paragraph_format.line_spacing = 1.5
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).width = Inches(1)
    table.cell(0, 1).width = Inches(2)
    doc.save(path)

    plan = load_hybrid_plan(path)
    exact_plan = plan.blocks[0]
    relative_plan = plan.blocks[1]
    table_plan = plan.blocks[2]
    assert exact_plan.left_indent == 72.0
    assert exact_plan.first_line_indent == 18.0
    assert exact_plan.line_spacing == 12.0
    assert relative_plan.line_spacing is None
    assert table_plan.column_widths == (72.0, 144.0)


def test_hybrid_parser_extracts_run_typography_and_ignores_theme_colour(tmp_path: Path) -> None:
    path = tmp_path / "run-typography.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    full = paragraph.add_run("Full")
    full.bold = True
    full.italic = True
    full.underline = True
    full.font.name = "Aptos"
    full.font.size = Pt(13)
    full.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    plain = paragraph.add_run("Plain")
    theme = paragraph.add_run("Theme")
    theme.font.color.theme_color = MSO_THEME_COLOR_INDEX.ACCENT_1
    doc.save(path)

    runs = load_hybrid_plan(path).blocks[0].runs
    assert runs[0].font_name == "Aptos"
    assert runs[0].font_size == 13.0
    assert runs[0].font_color == 0x123456
    assert runs[0].bold and runs[0].italic and runs[0].underline
    assert runs[1].font_name is None and runs[1].font_size is None and runs[1].font_color is None
    assert runs[2].font_color is None


def test_hybrid_parser_assigns_logical_list_groups(tmp_path: Path) -> None:
    path = tmp_path / "list-groups.docx"
    doc = Document()
    for text in ("One", "Two", "Three"):
        doc.add_paragraph(text, style="List Number")
    doc.add_paragraph("Normal")
    for text in ("Bullet one", "Bullet two"):
        doc.add_paragraph(text, style="List Bullet")
    doc.add_paragraph("Normal again")
    doc.add_paragraph("Later number", style="List Number")
    doc.save(path)

    plan = load_hybrid_plan(path)
    specs = [block.list_spec for block in plan.blocks]
    assert [(spec.kind, spec.group_id) if spec else None for spec in specs] == [
        ("number", 1), ("number", 1), ("number", 1), None,
        ("bullet", 2), ("bullet", 2), None, ("number", 3),
    ]


def test_hybrid_parser_isolates_list_groups_at_table_boundaries(tmp_path: Path) -> None:
    path = tmp_path / "table-list-boundary.docx"
    doc = Document()
    doc.add_paragraph("Before", style="List Number")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell"
    doc.add_paragraph("After", style="List Number")
    doc.save(path)

    plan = load_hybrid_plan(path)
    before_spec = plan.blocks[0].list_spec
    cell_spec = plan.blocks[1].rows[0][0].paragraphs[0].list_spec
    after_spec = plan.blocks[2].list_spec
    assert before_spec is not None and cell_spec is None and after_spec is not None
    assert before_spec.group_id != after_spec.group_id


def test_hybrid_parser_reports_merged_and_nested_tables(tmp_path: Path) -> None:
    path = tmp_path / "unsupported.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).add_table(rows=1, cols=1)
    doc.save(path)

    plan = load_hybrid_plan(path)
    assert any(item.kind == "merged table" for item in plan.unsupported)


def test_hybrid_dry_run_does_not_construct_com_adapter(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys) -> None:
    source = _docx(tmp_path)

    class ExplodingAdapter:
        def __init__(self) -> None:
            raise AssertionError("dry-run must not construct the COM adapter")

    monkeypatch.setattr(cli_module, "HybridWordAdapter", ExplodingAdapter)
    assert cli_module.main(["--target", "hybrid", "--dry-run", str(source)]) == 0
    output = capsys.readouterr().out
    assert "HYBRID DRY RUN" in output
    assert "COM: not imported or contacted" in output
    assert "Tables: 1" in output


class _FakeInformation:
    def __call__(self, index: int) -> bool:
        assert index == 12
        return False


class _FakeSelection:
    Start = End = 0
    StoryType = 1
    Information = _FakeInformation()

    def __init__(self, document) -> None:
        self.Document = document
        self.Range = self
        self.Font = _FakeFont()
        self.ParagraphFormat = type("Format", (), {})()
        self.ListFormat = _FakeListFormat()
        self.typed_paragraphs = 0

    def TypeParagraph(self) -> None:
        self.typed_paragraphs += 1


class _FakeListFormat:
    def __init__(self) -> None:
        self.operations: list[str] = []
        self.ListTemplate = "default-number-template"

    def ApplyBulletDefault(self) -> None:
        self.operations.append("bullet")

    def ApplyNumberDefault(self) -> None:
        self.operations.append("number")

    def ApplyListTemplate(self, template, continue_previous) -> None:
        self.operations.append(("restart", template, continue_previous))

    def RemoveNumbers(self) -> None:
        self.operations.append("remove")


class _FakeFont:
    def __init__(self) -> None:
        self.Bold = False
        self.Italic = False
        self.Underline = False
        self.Name = "destination-font"
        self.Size = 10.0
        self.Color = 0
        self.fail_property: str | None = None

    def __setattr__(self, name, value) -> None:
        if name != "fail_property" and getattr(self, "fail_property", None) == name:
            raise RuntimeError(f"failed {name}")
        object.__setattr__(self, name, value)


class _FakeDocument:
    ReadOnly = False
    ProtectionType = -1


class _FakeApp:
    def __init__(self) -> None:
        self.ActiveDocument = _FakeDocument()
        self.Selection = _FakeSelection(self.ActiveDocument)


def test_hybrid_adapter_rejects_unsafe_word_context() -> None:
    app = _FakeApp()
    app.ActiveDocument.ReadOnly = True
    with pytest.raises(HybridWordError, match="read-only"):
        HybridWordAdapter(active_object=lambda _: app).preflight()


@pytest.mark.parametrize(
    ("document", "story_type", "message"),
    [
        (None, 1, "no active document"),
        (_FakeDocument(), 6, "main document body"),
        (type("ProtectedDocument", (), {"ReadOnly": False, "ProtectionType": 0})(), 1, "protected"),
    ],
)
def test_hybrid_adapter_preflight_rejects_missing_protected_or_non_main_context(document, story_type, message) -> None:
    app = _FakeApp()
    app.ActiveDocument = document
    app.Selection.Document = document
    app.Selection.StoryType = story_type
    with pytest.raises(HybridWordError, match=message):
        HybridWordAdapter(active_object=lambda _: app).preflight()


def test_hybrid_adapter_explains_attach_failure() -> None:
    with pytest.raises(HybridWordError, match="will not launch"):
        HybridWordAdapter(active_object=lambda _: (_ for _ in ()).throw(OSError("missing"))).preflight()


def test_hybrid_adapter_applies_paragraph_and_run_formatting() -> None:
    app = _FakeApp()
    adapter = HybridWordAdapter(active_object=lambda _: app)
    context = adapter.preflight()
    paragraph = HybridParagraph((HybridRun("text", bold=True, italic=True, underline=True, font_name="Aptos", font_size=13.5, font_color=0x123456),), style_name="Heading 1", alignment=1, left_indent=10.0, line_spacing=12.0)
    adapter.prepare_paragraph(context, paragraph)
    adapter.prepare_run(context, paragraph.runs[0])
    assert app.Selection.Style == "Heading 1"
    assert app.Selection.ParagraphFormat.Alignment == 1
    assert app.Selection.ParagraphFormat.LeftIndent == 10.0
    assert app.Selection.Font.Bold is True
    assert app.Selection.Font.Italic is True
    assert app.Selection.Font.Underline is True
    assert app.Selection.Font.Name == "Aptos"
    assert app.Selection.Font.Size == 13.5
    assert app.Selection.Font.Color == 0x563412


def test_hybrid_adapter_preserves_unspecified_run_typography() -> None:
    app = _FakeApp()
    app.Selection.Font.Name = "keep-font"
    app.Selection.Font.Size = 17.0
    app.Selection.Font.Color = 0xABCDEF
    adapter = HybridWordAdapter(active_object=lambda _: app)
    context = adapter.preflight()

    adapter.prepare_run(context, HybridRun("text"))

    assert app.Selection.Font.Name == "keep-font"
    assert app.Selection.Font.Size == 17.0
    assert app.Selection.Font.Color == 0xABCDEF


@pytest.mark.parametrize("font_size", [0, -1, float("nan"), float("inf"), float("-inf")])
def test_hybrid_adapter_rejects_invalid_run_font_size(font_size: float) -> None:
    app = _FakeApp()
    adapter = HybridWordAdapter(active_object=lambda _: app)
    context = adapter.preflight()

    with pytest.raises(HybridWordError, match="invalid font size"):
        adapter.prepare_run(context, HybridRun("text", font_size=font_size))


def test_hybrid_adapter_reports_run_typography_com_failure() -> None:
    app = _FakeApp()
    app.Selection.Font.fail_property = "Color"
    adapter = HybridWordAdapter(active_object=lambda _: app)
    context = adapter.preflight()

    with pytest.raises(HybridWordError, match="applying hybrid run typography"):
        adapter.prepare_run(context, HybridRun("text", font_color=0xFF0000))


def test_hybrid_adapter_isolates_number_and_bullet_list_groups() -> None:
    app = _FakeApp()
    adapter = HybridWordAdapter(active_object=lambda _: app)
    context = adapter.preflight()

    paragraphs = (
        HybridParagraph((HybridRun("1"),), list_spec=HybridListSpec("number", 1)),
        HybridParagraph((HybridRun("2"),), list_spec=HybridListSpec("number", 1)),
        HybridParagraph((HybridRun("3"),), list_spec=HybridListSpec("number", 1)),
        HybridParagraph((HybridRun("normal"),)),
        HybridParagraph((HybridRun("bullet 1"),), list_spec=HybridListSpec("bullet", 2)),
        HybridParagraph((HybridRun("bullet 2"),), list_spec=HybridListSpec("bullet", 2)),
        HybridParagraph((HybridRun("normal again"),)),
        HybridParagraph((HybridRun("later number"),), list_spec=HybridListSpec("number", 3)),
    )
    for paragraph in paragraphs:
        adapter.prepare_paragraph(context, paragraph)

    assert app.Selection.ListFormat.operations == [
        "number",
        ("restart", "default-number-template", False),
        "remove",
        "bullet",
        "remove",
        "number",
        ("restart", "default-number-template", False),
    ]


def test_hybrid_adapter_reports_numbered_list_restart_failure() -> None:
    app = _FakeApp()

    def fail_restart(template, continue_previous) -> None:
        raise RuntimeError("restart failed")

    app.Selection.ListFormat.ApplyListTemplate = fail_restart
    adapter = HybridWordAdapter(active_object=lambda _: app)
    context = adapter.preflight()

    with pytest.raises(HybridWordError, match="restarting a hybrid numbered list at 1"):
        adapter.prepare_paragraph(
            context,
            HybridParagraph((HybridRun("numbered"),), list_spec=HybridListSpec("number", 1)),
        )


class _TableInsertionRange:
    def __init__(self, start: int = 10, end: int = 10) -> None:
        self.Start = start
        self.End = end
        self.collapse_direction = None
        self.selected = False

    @property
    def Duplicate(self):
        return _TableInsertionRange(self.Start, self.End)

    def Collapse(self, direction: int) -> None:
        self.collapse_direction = direction
        self.End = self.Start

    def Select(self) -> None:
        self.selected = True


class _TableSelection(_FakeSelection):
    def __init__(self, document) -> None:
        super().__init__(document)
        self.Range = _TableInsertionRange()


class _CreatedTable:
    def __init__(self) -> None:
        self.column_widths = {}
        self.Borders = type("Borders", (), {})()

        class Columns:
            def Item(columns_self, index: int):
                owner = self

                class Column:
                    @property
                    def Width(column_self):
                        return owner.column_widths.get(index)

                    @Width.setter
                    def Width(column_self, value: float) -> None:
                        owner.column_widths[index] = value

                return Column()

        self.Columns = Columns()


class _TableCollection:
    def __init__(self, failure: Exception | None = None) -> None:
        self.failure = failure
        self.received_range = None
        self.received_shape = None
        self.created = _CreatedTable()

    def Add(self, insertion_range, rows: int, columns: int):
        self.received_range = insertion_range
        self.received_shape = (rows, columns)
        if self.failure is not None:
            raise self.failure
        return self.created


class _TableDocument(_FakeDocument):
    def __init__(self, failure: Exception | None = None) -> None:
        self.Tables = _TableCollection(failure)


class _TableApp:
    def __init__(self, failure: Exception | None = None) -> None:
        self.ActiveDocument = _TableDocument(failure)
        self.Selection = _TableSelection(self.ActiveDocument)


def _simple_table() -> HybridTable:
    cell = HybridCell((HybridParagraph((HybridRun("cell"),)),))
    return HybridTable(((cell, cell),))


def test_hybrid_adapter_table_creation_uses_collapsed_duplicate_range() -> None:
    app = _TableApp()
    adapter = HybridWordAdapter(active_object=lambda _: app)
    context = adapter.preflight()

    created = adapter.create_table(context, _simple_table())

    insertion_range = app.ActiveDocument.Tables.received_range
    assert created is app.ActiveDocument.Tables.created
    assert app.ActiveDocument.Tables.received_shape == (1, 2)
    assert insertion_range is not app.Selection.Range
    assert insertion_range.collapse_direction == adapter._WD_COLLAPSE_END
    assert insertion_range.Start == insertion_range.End
    assert app.ActiveDocument.Tables.created.Borders.Enable is True


def test_hybrid_adapter_table_creation_surfaces_com_detail() -> None:
    app = _TableApp(RuntimeError("invalid insertion range"))
    adapter = HybridWordAdapter(active_object=lambda _: app)
    context = adapter.preflight()

    with pytest.raises(HybridWordError, match="invalid insertion range") as error:
        adapter.create_table(context, _simple_table())
    assert isinstance(error.value.__cause__, RuntimeError)


def test_hybrid_adapter_applies_point_widths_and_rejects_invalid_widths() -> None:
    app = _TableApp()
    adapter = HybridWordAdapter(active_object=lambda _: app)
    context = adapter.preflight()

    adapter.create_table(context, HybridTable(_simple_table().rows, (72.0, 144.0)))
    assert app.ActiveDocument.Tables.created.column_widths == {1: 72.0, 2: 144.0}

    failing_app = _TableApp()
    failing_adapter = HybridWordAdapter(active_object=lambda _: failing_app)
    failing_context = failing_adapter.preflight()
    with pytest.raises(HybridWordError, match="invalid width"):
        failing_adapter.create_table(failing_context, HybridTable(_simple_table().rows, (72.0, 2000.0)))
    assert failing_app.ActiveDocument.Tables.received_range is None

    class BrokenBorders:
        @property
        def Enable(self):
            return False

        @Enable.setter
        def Enable(self, value):
            raise RuntimeError("border operation failed")

    border_app = _TableApp()
    border_app.ActiveDocument.Tables.created.Borders = BrokenBorders()
    border_adapter = HybridWordAdapter(active_object=lambda _: border_app)
    border_context = border_adapter.preflight()
    with pytest.raises(HybridWordError, match="applying the hybrid table grid: border operation failed"):
        border_adapter.create_table(border_context, _simple_table())


def test_hybrid_adapter_cell_target_excludes_end_marker() -> None:
    app = _TableApp()
    adapter = HybridWordAdapter(active_object=lambda _: app)
    context = adapter.preflight()

    class CellRange(_TableInsertionRange):
        def __init__(self, start: int, end: int) -> None:
            super().__init__(start, end)
            self.pre_collapse_end = None

        @property
        def Duplicate(self):
            return self

        def Collapse(self, direction: int) -> None:
            self.pre_collapse_end = self.End
            super().Collapse(direction)

    cell_range = CellRange(20, 27)
    cell_range.ListFormat = app.Selection.ListFormat
    app.Selection.Range = cell_range

    class Cell:
        Range = cell_range

    class Table:
        @staticmethod
        def Cell(row: int, column: int):
            assert (row, column) == (1, 1)
            return Cell()

    adapter.prepare_cell(context, Table(), 1, 1, _simple_table().rows[0][0])

    assert cell_range.pre_collapse_end == 26
    assert cell_range.collapse_direction == adapter._WD_COLLAPSE_END
    assert cell_range.selected is True


class _RecordingAdapter:
    def __init__(self, events: list[str]) -> None:
        self.events = events

    def preflight(self):
        self.events.append("preflight")
        return object()

    def prepare_paragraph(self, context, paragraph) -> None:
        self.events.append("paragraph")

    def prepare_run(self, context, run) -> None:
        self.events.append(f"format:{run.text}")

    def advance_paragraph(self, context) -> None:
        self.events.append("advance")

    def create_table(self, context, table):
        self.events.append("table")
        return object()

    def prepare_cell(self, context, table, row, column, cell) -> None:
        self.events.append(f"cell:{row},{column}")

    def move_after_table(self, context, table) -> None:
        self.events.append("after-table")


class _RecordingExecutor:
    def __init__(self, events: list[str]) -> None:
        self.events = events

    def type_text(self, text: str) -> None:
        self.events.append(f"type:{text}")

    def press_key(self, key: str) -> None:
        self.events.append(f"key:{key}")


class _NoopMonitor:
    def __init__(self, **kwargs) -> None:
        self.kwargs = kwargs

    def start(self) -> None:
        return None

    def close(self) -> None:
        return None


class _HandoffController:
    def __init__(self, executor, status_callback, events) -> None:
        self._executor = executor
        self._status_callback = status_callback
        self._events = events

    def run(self, actions, countdown_seconds=None):
        if countdown_seconds:
            self._events.append("handoff-countdown")
            self._status_callback(RunState.COUNTDOWN, "Starting in 5 seconds...")
            return RunResult(RunState.IDLE, 0, 0)
        self._events.append("typing-controller")
        typed = 0
        for action in actions:
            if isinstance(action, TypeText):
                for character in action.text:
                    self._executor.type_text(character)
                    typed += 1
        return RunResult(RunState.IDLE, len(tuple(actions)), typed)


def test_hybrid_runner_orders_structure_format_focus_and_typing() -> None:
    events: list[str] = []
    plan = HybridDocumentPlan(Path("source.docx"), (HybridParagraph((HybridRun("Hi", bold=True),)),))
    runner = HybridRunner(
        adapter=_RecordingAdapter(events), focus_guard=WordFocusGuard(lambda: events.append("focus") is None),
        executor=_RecordingExecutor(events), typing_config=TypingConfig(words_per_minute=1000, countdown_seconds=0),
        profile="precise", seed=1, hotkey_monitor_factory=_NoopMonitor, status=lambda _: None,
    )
    result = runner.run(plan)
    assert result.completed_targets == 1
    assert events[:4] == ["preflight", "paragraph", "format:Hi", "focus"]
    assert events[-2:] == ["type:H", "type:i"]


def test_hybrid_runner_handoff_precedes_first_focus_check_and_typing() -> None:
    events: list[str] = []
    plan = HybridDocumentPlan(Path("source.docx"), (HybridParagraph((HybridRun("Hi"),)),))

    def controller_factory(*, executor, config, status_callback):
        return _HandoffController(executor, status_callback, events)

    def status(message: str) -> None:
        if message.startswith("[Hybrid mode] Switch"):
            events.append("handoff-instruction")
        elif message.startswith("[Hybrid mode] Typing"):
            events.append("typing-status")

    runner = HybridRunner(
        adapter=_RecordingAdapter(events), focus_guard=WordFocusGuard(lambda: events.append("focus") is None),
        executor=_RecordingExecutor(events), typing_config=TypingConfig(words_per_minute=1000, countdown_seconds=5),
        profile="precise", seed=1, hotkey_monitor_factory=_NoopMonitor, controller_factory=controller_factory,
        status=status,
    )
    runner.run(plan)

    assert events[:8] == [
        "preflight", "paragraph", "format:Hi", "handoff-instruction", "handoff-countdown",
        "focus", "typing-status", "typing-controller",
    ]
    assert events[-2:] == ["type:H", "type:i"]


def test_hybrid_runner_failed_focus_after_handoff_does_not_type() -> None:
    events: list[str] = []
    plan = HybridDocumentPlan(Path("source.docx"), (HybridParagraph((HybridRun("Hi"),)),))

    def controller_factory(*, executor, config, status_callback):
        return _HandoffController(executor, status_callback, events)

    runner = HybridRunner(
        adapter=_RecordingAdapter(events), focus_guard=WordFocusGuard(lambda: False),
        executor=_RecordingExecutor(events), typing_config=TypingConfig(words_per_minute=1000, countdown_seconds=5),
        profile="precise", seed=None, hotkey_monitor_factory=_NoopMonitor, controller_factory=controller_factory,
        status=lambda message: events.append(message),
    )
    with pytest.raises(HybridRunError, match="partially modified"):
        runner.run(plan)

    assert "handoff-countdown" in events
    assert not any(item.startswith("type:") for item in events)


def test_hybrid_runner_keeps_focus_check_for_later_segments() -> None:
    events: list[str] = []
    plan = HybridDocumentPlan(
        Path("source.docx"),
        (HybridParagraph((HybridRun("A"), HybridRun("B"))),),
    )

    def controller_factory(*, executor, config, status_callback):
        return _HandoffController(executor, status_callback, events)

    runner = HybridRunner(
        adapter=_RecordingAdapter(events), focus_guard=WordFocusGuard(lambda: events.append("focus") is None),
        executor=_RecordingExecutor(events), typing_config=TypingConfig(words_per_minute=1000, countdown_seconds=5),
        profile="precise", seed=1, hotkey_monitor_factory=_NoopMonitor, controller_factory=controller_factory,
        status=lambda _: None,
    )
    runner.run(plan)

    assert events.count("handoff-countdown") == 1
    assert events.count("focus") == 2
    assert [item for item in events if item.startswith("type:")] == ["type:A", "type:B"]


def test_hybrid_runner_preserves_paragraph_table_paragraph_transition() -> None:
    events: list[str] = []
    plan = HybridDocumentPlan(
        Path("source.docx"),
        (
            HybridParagraph((HybridRun("before"),)),
            _simple_table(),
            HybridParagraph((HybridRun("after"),)),
        ),
    )
    runner = HybridRunner(
        adapter=_RecordingAdapter(events), focus_guard=WordFocusGuard(lambda: True),
        executor=_RecordingExecutor(events), typing_config=TypingConfig(words_per_minute=1000, countdown_seconds=0),
        profile="precise", seed=1, hotkey_monitor_factory=_NoopMonitor, status=lambda _: None,
    )

    runner.run(plan)

    before = events.index("format:before")
    advance = events.index("advance")
    table = events.index("table")
    cell = events.index("cell:1,1")
    after_table = events.index("after-table")
    after = events.index("format:after")
    assert before < advance < table < cell < after_table < after


def test_hybrid_runner_stops_before_typing_after_focus_loss() -> None:
    events: list[str] = []
    plan = HybridDocumentPlan(Path("source.docx"), (HybridParagraph((HybridRun("Hi"),)),))
    runner = HybridRunner(
        adapter=_RecordingAdapter(events), focus_guard=WordFocusGuard(lambda: False),
        executor=_RecordingExecutor(events), typing_config=TypingConfig(words_per_minute=1000, countdown_seconds=0),
        profile="precise", seed=None, hotkey_monitor_factory=_NoopMonitor, status=lambda _: None,
    )
    with pytest.raises(HybridRunError, match="partially modified"):
        runner.run(plan)
    assert not any(item.startswith("type:") for item in events)


def test_hybrid_runner_rejects_unsupported_before_com_mutation() -> None:
    plan = HybridDocumentPlan(Path("source.docx"), (), (HybridUnsupported("image", "unsupported"),))
    events: list[str] = []
    runner = HybridRunner(
        adapter=_RecordingAdapter(events), focus_guard=WordFocusGuard(lambda: True), executor=_RecordingExecutor(events),
        typing_config=TypingConfig(), profile="precise", seed=None, hotkey_monitor_factory=_NoopMonitor, status=lambda _: None,
    )
    with pytest.raises(HybridRunError, match="cannot safely"):
        runner.run(plan)
    assert events == []


def test_hybrid_runner_reports_preflight_failure_without_typing() -> None:
    class FailingAdapter(_RecordingAdapter):
        def preflight(self):
            raise HybridWordError("no active document")

    events: list[str] = []
    runner = HybridRunner(
        adapter=FailingAdapter(events), focus_guard=WordFocusGuard(lambda: True), executor=_RecordingExecutor(events),
        typing_config=TypingConfig(), profile="precise", seed=None, hotkey_monitor_factory=_NoopMonitor, status=lambda _: None,
    )
    with pytest.raises(HybridRunError, match="no active document"):
        runner.run(HybridDocumentPlan(Path("source.docx"), (HybridParagraph((HybridRun("text"),)),)))
    assert not any(item.startswith("type:") for item in events)
